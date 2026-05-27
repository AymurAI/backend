import base64
import json
import re
import subprocess
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path
from unittest.mock import MagicMock, patch

import pymupdf
import pytest
from docx import Document

from aymurai.database.schema import AnonymizationParagraph
from aymurai.database.utils import text_to_uuid
from aymurai.meta.api_interfaces import LabelPolicy, RenderPolicy
from aymurai.text.anonymization import DocxAnonymizer, PdfAnonymizer, get_anonymizer
from aymurai.text.anonymization.alignment import index_paragraphs
from aymurai.text.anonymization.pdf.ops import _refine_signature_text_rect
from aymurai.text.anonymization.pdf.widgets import _signature_background_rect
from tests.api.conftest import build_label
from tests.api.routers.conftest import build_mock_pipeline

PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+a6R8AAAAASUVORK5CYII="
)
PNG_BLACK_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAAAAAA6fptVAAAACklEQVR42mNgAAAAAgAB4iG8MwAAAABJRU5ErkJggg=="
)
WATERMARK_URL = "https://www.aymurai.info/"

WINDOWS_PYMUPDF_LAYOUT_XFAIL = pytest.mark.xfail(
    sys.platform == "win32",
    reason="pymupdf4llm ONNX layout model receives int32 tensors on Windows (expects int64)",
    strict=False,
)


def _write_pdf(path: Path, configure) -> Path:
    doc = pymupdf.open()
    page = doc.new_page()
    configure(doc, page)
    doc.save(path)
    doc.close()
    return path


def _label_dict(text: str, label: str = "PER", **attrs) -> dict:
    payload = build_label(label, text).model_dump(mode="json")
    payload["attrs"].update(attrs)
    return payload


def _run_pdf_anonymizer(
    tmp_path: Path,
    source_path: Path,
    document: str,
    labels: list[dict],
    render_context: dict | None = None,
) -> Path:
    output_dir = tmp_path / "out"
    output_dir.mkdir(exist_ok=True)
    output_path = PdfAnonymizer().anonymize(
        {"path": str(source_path)},
        [{"document": document, "labels": labels}],
        str(output_dir),
        render_context=render_context,
    )
    return Path(output_path)


def _label_for_document_text(document: str, text: str, label: str = "PER") -> dict:
    payload = _label_dict(text, label)
    start = document.index(text)
    payload["start_char"] = start
    payload["end_char"] = start + len(text)
    return payload


def _render_context_for_entities(labels: list[dict]) -> dict:
    index_by_entity = {}
    next_index_by_base = {}
    for label in labels:
        attrs = label.get("attrs") or {}
        base = attrs.get("aymurai_label") or label.get("label") or "ENT"
        entity_id = str(attrs.get("canonical_entity_id") or label.get("text"))
        key = (base, entity_id)
        if key not in index_by_entity:
            next_index_by_base[base] = next_index_by_base.get(base, 0) + 1
            index_by_entity[key] = next_index_by_base[base]

    return {
        "render_policy": RenderPolicy(suffix_mode="always", suffix_threshold=0),
        "label_policies": {"PER": LabelPolicy()},
        "count_by_base": dict(next_index_by_base),
        "index_by_entity": index_by_entity,
    }


def _dark_pixel_ratio(page: pymupdf.Page, rect: pymupdf.Rect) -> float:
    pixmap = page.get_pixmap(
        matrix=pymupdf.Matrix(2, 2),
        clip=rect,
        alpha=False,
    )
    samples = pixmap.samples
    if not samples:
        return 0.0

    channels = pixmap.n
    dark_pixels = 0
    total_pixels = pixmap.width * pixmap.height
    for offset in range(0, len(samples), channels):
        if all(channel < 96 for channel in samples[offset : offset + 3]):
            dark_pixels += 1

    return dark_pixels / max(total_pixels, 1)


def _assert_text_count(page_text: str, text: str, expected: int) -> None:
    assert page_text.count(text) == expected, page_text


def _assert_rect_close(actual: pymupdf.Rect, expected: pymupdf.Rect) -> None:
    assert (actual.x0, actual.y0, actual.x1, actual.y1) == pytest.approx(
        (expected.x0, expected.y0, expected.x1, expected.y1)
    )


def _write_variable_signature_pdf(
    path: Path,
) -> tuple[Path, list[dict], list[str], list[str], list[pymupdf.Rect]]:
    blocks = [
        {
            "origin": (58, 112),
            "lines": [
                "Mesa de Control 42",
                "Adriana Morales",
                "Area de Validacion",
                "Codigo A-17",
            ],
            "signer": "Adriana Morales",
            "qr": "top",
        },
        {
            "origin": (326, 112),
            "lines": [
                "Bernardo Diaz",
                "Direccion Legal",
                "Organismo Beta Sur",
                "Tramite BX-900",
            ],
            "signer": "Bernardo Diaz",
            "qr": "right",
        },
        {
            "origin": (58, 328),
            "lines": [
                "Centro Operativo",
                "Carolina Ruiz",
                "Secretaria Tecnica",
                "2026-05-26 10:15",
            ],
            "signer": "Carolina Ruiz",
            "qr": "left",
        },
        {
            "origin": (326, 328),
            "lines": [
                "Unidad Regional",
                "Coordinacion de Revision",
                "Daniel Silva",
                "Expediente Digital Z-42",
            ],
            "signer": "Daniel Silva",
            "qr": "top",
        },
        {
            "origin": (58, 544),
            "lines": [
                "Responsable: Elena Torres - Acta Final",
                "Delegacion Gamma",
                "Registro Interno R-204",
            ],
            "signer": "Elena Torres",
            "qr": "right",
        },
    ]

    doc = pymupdf.open()
    page = doc.new_page()
    preds: list[dict] = []
    preserved_texts: list[str] = []
    signers: list[str] = []
    qr_rects: list[pymupdf.Rect] = []

    for idx, block in enumerate(blocks):
        x, y = block["origin"]
        if block["qr"] == "right":
            qr_rect = pymupdf.Rect(x + 150, y - 4, x + 182, y + 28)
        elif block["qr"] == "left":
            qr_rect = pymupdf.Rect(x - 2, y - 48, x + 30, y - 16)
        else:
            qr_rect = pymupdf.Rect(x, y - 52, x + 32, y - 20)
        page.insert_image(qr_rect, stream=PNG_BLACK_1X1)
        qr_rects.append(qr_rect)

        for line_idx, line in enumerate(block["lines"]):
            page.insert_text((x, y + (line_idx * 16)), line, fontsize=11)
            if line == block["signer"]:
                continue
            if block["signer"] in line:
                preserved_texts.extend(
                    part.strip() for part in line.split(block["signer"]) if part.strip()
                )
            else:
                preserved_texts.append(line)

        widget = pymupdf.Widget()
        widget.field_name = f"sig_{idx}"
        widget.field_type = pymupdf.PDF_WIDGET_TYPE_SIGNATURE
        widget.rect = pymupdf.Rect(x - 12, y - 62, x + 230, y + 64)
        page.add_widget(widget)

        document = "\n".join(block["lines"])
        label = _label_for_document_text(document, block["signer"])
        preds.append({"document": document, "labels": [label]})
        signers.append(block["signer"])

    doc.save(path)
    doc.close()
    return path, preds, signers, preserved_texts, qr_rects


@pytest.mark.integration
def test_anonymization_package_exports_and_registry_are_stable():
    assert PdfAnonymizer.__name__ == "PdfAnonymizer"
    assert DocxAnonymizer.__name__ == "DocxAnonymizer"
    assert isinstance(get_anonymizer("pdf"), PdfAnonymizer)
    assert isinstance(get_anonymizer("docx"), DocxAnonymizer)


@pytest.mark.integration
@WINDOWS_PYMUPDF_LAYOUT_XFAIL
def test_pdf_anonymizer_falls_back_from_invalid_alt_offsets(tmp_path):
    document = "Ana Perez firmo el escrito"
    source_path = _write_pdf(
        tmp_path / "invalid-alt.pdf",
        lambda _doc, page: page.insert_text((72, 72), document),
    )
    labels = [
        _label_dict(
            "Ana Perez",
            aymurai_alt_start_char=999,
            aymurai_alt_end_char=1000,
        )
    ]

    output_path = _run_pdf_anonymizer(tmp_path, source_path, document, labels)

    with pymupdf.open(output_path) as output_doc:
        page_text = output_doc[0].get_text()

    assert "Ana Perez" not in page_text
    assert "<PER>" in page_text


@pytest.mark.integration
@WINDOWS_PYMUPDF_LAYOUT_XFAIL
def test_pdf_anonymizer_scrubs_pdf_payloads_and_preserves_safe_links(tmp_path):
    document = "Ana Perez presento el escrito"

    def configure(doc: pymupdf.Document, page: pymupdf.Page) -> None:
        page.insert_text((72, 72), document)
        sensitive_rect = page.search_for("Ana Perez")[0]
        page.insert_link(
            {
                "kind": pymupdf.LINK_URI,
                "from": sensitive_rect,
                "uri": "https://secret.example",
            }
        )
        safe_rect = pymupdf.Rect(72, 140, 180, 155)
        page.insert_text((72, 150), "Portal publico")
        page.insert_link(
            {
                "kind": pymupdf.LINK_URI,
                "from": safe_rect,
                "uri": "https://safe.example",
            }
        )
        page.add_file_annot((220, 72), b"attached secret", "attached.txt")
        doc.set_metadata(
            {
                "title": "Secret title",
                "author": "Secret author",
                "subject": "Secret subject",
                "keywords": "alpha,beta",
                "creator": "Secret creator",
                "producer": "Secret producer",
            }
        )
        doc.set_xml_metadata("<x:xmpmeta>top-secret</x:xmpmeta>")
        doc.embfile_add("secret.txt", b"secret bytes", filename="secret.txt")

    source_path = _write_pdf(tmp_path / "metadata.pdf", configure)
    labels = [_label_dict("Ana Perez")]

    output_path = _run_pdf_anonymizer(tmp_path, source_path, document, labels)

    with pymupdf.open(output_path) as output_doc:
        page = output_doc[0]
        link_uris = {link.get("uri") for link in page.get_links()}

        assert output_doc.metadata.get("title") == ""
        assert output_doc.metadata.get("subject") == ""
        assert output_doc.metadata.get("keywords") == ""
        assert output_doc.metadata.get("creationDate") == ""
        assert re.fullmatch(
            r"D:\d{14}\+00'00'",
            output_doc.metadata.get("modDate") or "",
        )
        assert output_doc.metadata.get("trapped") == ""
        assert output_doc.metadata.get("author") == ""
        assert output_doc.metadata.get("creator") == "AymurAI"
        assert output_doc.metadata.get("producer") == "AymurAI"
        assert not output_doc.get_xml_metadata()
        assert output_doc.embfile_names() == []
        assert list(page.annots() or []) == []
        assert "https://secret.example" not in link_uris
        assert "https://safe.example" in link_uris
        assert WATERMARK_URL in link_uris


@pytest.mark.integration
@WINDOWS_PYMUPDF_LAYOUT_XFAIL
def test_pdf_anonymizer_moves_watermark_away_from_footer_content(tmp_path):
    document = "Ana Perez presento el escrito"
    footer_rect = pymupdf.Rect(360, 760, 575, 815)

    def configure(_doc: pymupdf.Document, page: pymupdf.Page) -> None:
        page.insert_text((72, 72), document)
        page.draw_rect(footer_rect, color=(0, 0, 0), fill=(0, 0, 0), overlay=True)

    source_path = _write_pdf(tmp_path / "footer-watermark.pdf", configure)
    output_path = _run_pdf_anonymizer(
        tmp_path,
        source_path,
        document,
        [_label_dict("Ana Perez")],
    )

    with pymupdf.open(output_path) as output_doc:
        page = output_doc[0]
        watermark_links = [
            link for link in page.get_links() if link.get("uri") == WATERMARK_URL
        ]

        assert len(watermark_links) == 1
        watermark_rect = pymupdf.Rect(watermark_links[0]["from"])
        assert not watermark_rect.intersects(footer_rect)
        assert watermark_rect.x1 < footer_rect.x0


@pytest.mark.integration
@WINDOWS_PYMUPDF_LAYOUT_XFAIL
def test_pdf_anonymizer_removes_image_backed_entities(tmp_path):
    source_path = _write_pdf(
        tmp_path / "image.pdf",
        lambda _doc, page: (
            page.insert_image(pymupdf.Rect(60, 60, 220, 110), stream=PNG_1X1),
            page.insert_text((80, 90), "Ana Perez"),
        ),
    )

    output_path = _run_pdf_anonymizer(
        tmp_path,
        source_path,
        "Ana Perez",
        [_label_dict("Ana Perez")],
    )

    with pymupdf.open(output_path) as output_doc:
        page = output_doc[0]
        page_text = page.get_text()

        assert page.get_image_info() == []
        assert "Ana Perez" not in page_text
        assert "<PER>" in page_text


def test_signature_background_rect_stays_on_signer_name_line():
    background = _signature_background_rect(
        {
            "line_rect": pymupdf.Rect(80, 70, 220, 118),
            "canvas_rect": pymupdf.Rect(112, 70, 145, 82),
            "redact_rect": pymupdf.Rect(112, 70, 145, 82),
        },
        pymupdf.Rect(60, 60, 230, 130),
    )

    assert background.y0 >= 70
    assert background.y1 <= 82


def test_signature_text_rect_refinement_does_not_include_role_text(tmp_path):
    source_path = _write_pdf(
        tmp_path / "signature-role.pdf",
        lambda _doc, page: (
            page.insert_text((100, 80), "RUIZ"),
            page.insert_text((100, 96), "JUEZ/A"),
        ),
    )

    with pymupdf.open(source_path) as doc:
        page = doc[0]
        signer_rect = page.search_for("RUIZ")[0]
        role_rect = page.search_for("JUEZ/A")[0]
        loose_rect = pymupdf.Rect(signer_rect)
        loose_rect.include_rect(role_rect)

        refined = _refine_signature_text_rect(
            page,
            "RUIZ",
            pymupdf.Rect(80, 60, 200, 115),
            loose_rect,
        )

    assert refined.intersects(signer_rect)
    assert not refined.intersects(role_rect)


def test_signature_text_rect_refinement_returns_current_rect_when_no_hit_in_widget(
    tmp_path,
):
    source_path = _write_pdf(
        tmp_path / "signature-no-hit.pdf",
        lambda _doc, page: page.insert_text((260, 80), "RUIZ"),
    )

    with pymupdf.open(source_path) as doc:
        page = doc[0]
        current_rect = pymupdf.Rect(100, 72, 130, 84)

        refined = _refine_signature_text_rect(
            page,
            "RUIZ",
            pymupdf.Rect(80, 60, 180, 115),
            current_rect,
        )

    _assert_rect_close(refined, current_rect)


def test_signature_text_rect_refinement_selects_closest_matching_hit(tmp_path):
    source_path = _write_pdf(
        tmp_path / "signature-multiple-hits.pdf",
        lambda _doc, page: (
            page.insert_text((100, 80), "RUIZ"),
            page.insert_text((220, 80), "RUIZ"),
        ),
    )

    with pymupdf.open(source_path) as doc:
        page = doc[0]
        left_rect, right_rect = page.search_for("RUIZ")
        target = pymupdf.Rect(right_rect)
        target.x0 += 2
        target.x1 += 2

        refined = _refine_signature_text_rect(
            page,
            "RUIZ",
            pymupdf.Rect(80, 60, 280, 115),
            target,
        )

    assert refined.intersects(right_rect)
    assert not refined.intersects(left_rect)


@pytest.mark.integration
@WINDOWS_PYMUPDF_LAYOUT_XFAIL
def test_pdf_anonymizer_only_redacts_marked_signature_names_in_variable_layouts(
    tmp_path,
):
    source_path, preds, signers, preserved_texts, qr_rects = (
        _write_variable_signature_pdf(tmp_path / "variable-signatures.pdf")
    )
    render_context = _render_context_for_entities([pred["labels"][0] for pred in preds])
    output_dir = tmp_path / "out-variable"
    output_dir.mkdir(exist_ok=True)

    output_path = PdfAnonymizer().anonymize(
        {"path": str(source_path)},
        preds,
        str(output_dir),
        render_context=render_context,
    )

    with pymupdf.open(output_path) as output_doc:
        page = output_doc[0]
        page_text = page.get_text()

        assert list(page.widgets() or []) == []
        assert len(page.get_image_info()) >= len(qr_rects)

        for signer in signers:
            assert signer not in page_text

        for index in range(1, len(signers) + 1):
            assert f"<PER_{index}>" in page_text

        for preserved_text in preserved_texts:
            _assert_text_count(page_text, preserved_text, 1)

        for qr_rect in qr_rects:
            assert _dark_pixel_ratio(page, qr_rect) > 0.25


@pytest.mark.integration
@WINDOWS_PYMUPDF_LAYOUT_XFAIL
def test_pdf_anonymizer_leaves_unlabeled_signature_names_visible(tmp_path):
    source_path, preds, signers, preserved_texts, qr_rects = (
        _write_variable_signature_pdf(tmp_path / "partially-labeled-signatures.pdf")
    )
    unlabeled_signer = signers[-1]
    filtered_preds = []
    filtered_labels = []
    for pred, signer in zip(preds, signers, strict=True):
        labels = [] if signer == unlabeled_signer else pred["labels"]
        filtered_preds.append({**pred, "labels": labels})
        filtered_labels.extend(labels)

    render_context = _render_context_for_entities(filtered_labels)
    output_dir = tmp_path / "out-partial"
    output_dir.mkdir(exist_ok=True)

    output_path = PdfAnonymizer().anonymize(
        {"path": str(source_path)},
        filtered_preds,
        str(output_dir),
        render_context=render_context,
    )

    with pymupdf.open(output_path) as output_doc:
        page = output_doc[0]
        page_text = page.get_text()

        assert list(page.widgets() or []) == []
        assert unlabeled_signer in page_text
        assert "<PER_5>" not in page_text

        for index, signer in enumerate(signers[:-1], start=1):
            assert signer not in page_text
            assert f"<PER_{index}>" in page_text

        for preserved_text in preserved_texts:
            _assert_text_count(page_text, preserved_text, 1)

        for qr_rect in qr_rects:
            assert _dark_pixel_ratio(page, qr_rect) > 0.25


@pytest.mark.integration
@WINDOWS_PYMUPDF_LAYOUT_XFAIL
def test_pdf_anonymizer_preserves_non_signature_widget_appearance_when_baking(
    tmp_path,
):
    def configure(_doc: pymupdf.Document, page: pymupdf.Page) -> None:
        page.insert_text((80, 88), "Ana Perez")

        text_widget = pymupdf.Widget()
        text_widget.field_name = "public_field"
        text_widget.field_type = pymupdf.PDF_WIDGET_TYPE_TEXT
        text_widget.field_value = "Visible Field Value"
        text_widget.text_font = "Helv"
        text_widget.text_fontsize = 10
        text_widget.rect = pymupdf.Rect(260, 70, 410, 96)
        page.add_widget(text_widget)

        signature_widget = pymupdf.Widget()
        signature_widget.field_name = "sig_1"
        signature_widget.field_type = pymupdf.PDF_WIDGET_TYPE_SIGNATURE
        signature_widget.rect = pymupdf.Rect(60, 60, 180, 110)
        page.add_widget(signature_widget)

    source_path = _write_pdf(tmp_path / "signature-and-text-widget.pdf", configure)
    output_path = _run_pdf_anonymizer(
        tmp_path,
        source_path,
        "Ana Perez",
        [_label_dict("Ana Perez")],
    )

    with pymupdf.open(output_path) as output_doc:
        page = output_doc[0]
        page_text = page.get_text()

        assert list(page.widgets() or []) == []
        assert "Visible Field Value" in page_text
        assert "Ana Perez" not in page_text
        assert "<PER>" in page_text


@pytest.mark.integration
@WINDOWS_PYMUPDF_LAYOUT_XFAIL
def test_pdf_anonymizer_preserves_signature_appearance_when_redacting_signer_name(
    tmp_path,
):
    def configure(_doc: pymupdf.Document, page: pymupdf.Page) -> None:
        page.insert_text((80, 76), "FIRMADO DIGITALMENTE")
        page.insert_text((80, 92), "05/02/2025 14:17")
        page.insert_text((80, 108), "Ana Perez")
        page.insert_image(pymupdf.Rect(185, 68, 215, 98), stream=PNG_1X1)
        widget = pymupdf.Widget()
        widget.field_name = "sig_1"
        widget.field_type = pymupdf.PDF_WIDGET_TYPE_SIGNATURE
        widget.rect = pymupdf.Rect(60, 60, 230, 120)
        page.add_widget(widget)

    source_path = _write_pdf(tmp_path / "signature.pdf", configure)
    output_path = _run_pdf_anonymizer(
        tmp_path,
        source_path,
        "Ana Perez",
        [_label_dict("Ana Perez")],
    )

    with pymupdf.open(output_path) as output_doc:
        page = output_doc[0]
        page_text = page.get_text()

        assert list(page.widgets() or []) == []
        assert page.get_image_info() != []
        assert "FIRMADO DIGITALMENTE" in page_text
        assert "05/02/2025 14:17" in page_text
        assert "Ana Perez" not in page_text
        assert "<PER>" in page_text


def test_index_paragraphs_reads_docx_xml_as_utf8(tmp_path):
    xml_path = tmp_path / "document.xml"
    xml_path.write_bytes(
        """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p><w:r><w:t>Señora — resolución</w:t></w:r></w:p>
  </w:body>
</w:document>
""".encode("utf-8")
    )

    paragraphs = index_paragraphs(str(xml_path))

    assert len(paragraphs) == 1
    assert paragraphs[0]["plain_text"] == "Señora — resolución"


@pytest.mark.integration
def test_docx_anonymizer_sets_aymurai_core_properties(tmp_path):
    source_path = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("Ana Perez firmo el escrito")
    document.core_properties.author = "Sensitive Author"
    document.core_properties.last_modified_by = "Sensitive Modifier"
    document.save(source_path)

    started_at = datetime.now(timezone.utc).replace(microsecond=0)

    output_path = DocxAnonymizer().anonymize(
        {"path": str(source_path)},
        [
            {
                "document": "Ana Perez firmo el escrito",
                "labels": [_label_dict("Ana Perez")],
            }
        ],
        str(tmp_path / "out"),
    )

    output_document = Document(output_path)
    core_properties = output_document.core_properties
    assert core_properties.author == ""
    assert core_properties.last_modified_by == "AymurAI"
    assert core_properties.modified is not None
    modified = core_properties.modified
    if modified.tzinfo is None:
        modified = modified.replace(tzinfo=timezone.utc)
    assert started_at <= modified <= datetime.now(timezone.utc) + timedelta(seconds=5)


@pytest.mark.integration
@patch("aymurai.api.endpoints.routers.anonymizer.anonymizer.load_pipeline")
def test_should_return_prediction_when_text_provided(mock_load_pipeline, client):
    mock_pipeline = build_mock_pipeline()
    mock_load_pipeline.return_value = mock_pipeline

    response = client.post(
        "/anonymizer/predict",
        json={"text": "Sample anonymization text"},
    )

    assert response.status_code == 200
    data = response.json()
    assert "document" in data
    assert "labels" in data
    assert data["document"] == "Sample anonymization text"
    assert isinstance(data["labels"], list)


@pytest.mark.integration
@patch("aymurai.api.endpoints.routers.anonymizer.anonymizer.load_pipeline")
def test_should_return_cached_prediction_when_text_in_cache(
    mock_load_pipeline, client, db_session
):
    text = "Cached text with entities"
    labels = [build_label("PER", "Juan Pérez").model_dump(mode="json")]

    paragraph_id = text_to_uuid(text)
    cached_para = AnonymizationParagraph(
        id=paragraph_id,
        text=text,
        prediction=labels,
    )
    db_session.add(cached_para)
    db_session.commit()

    response = client.post(
        "/anonymizer/predict",
        json={"text": text},
        params={"use_cache": True},
    )

    assert response.status_code == 200
    data = response.json()
    assert data["document"] == text
    assert data["labels"] == labels
    mock_load_pipeline.assert_not_called()


@pytest.mark.integration
@patch("aymurai.api.endpoints.routers.anonymizer.anonymizer.load_pipeline")
def test_should_store_prediction_in_db_when_use_cache_true(
    mock_load_pipeline, client, db_session
):
    mock_pipeline = build_mock_pipeline()
    mock_load_pipeline.return_value = mock_pipeline

    text = "New prediction to cache"
    response = client.post(
        "/anonymizer/predict",
        json={"text": text},
        params={"use_cache": True},
    )

    assert response.status_code == 200

    paragraph_id = text_to_uuid(text)
    stored = db_session.get(AnonymizationParagraph, paragraph_id)
    assert stored is not None
    assert stored.text == text
    assert stored.prediction is not None


@pytest.mark.integration
@patch("aymurai.api.endpoints.routers.anonymizer.anonymizer.load_pipeline")
def test_should_return_cached_result_when_calling_twice(mock_load_pipeline, client):
    mock_pipeline = build_mock_pipeline()
    mock_load_pipeline.return_value = mock_pipeline

    text = "Repeated query text"

    response1 = client.post(
        "/anonymizer/predict",
        json={"text": text},
        params={"use_cache": True},
    )
    data1 = response1.json()

    response2 = client.post(
        "/anonymizer/predict",
        json={"text": text},
        params={"use_cache": True},
    )
    data2 = response2.json()

    assert response1.status_code == 200
    assert response2.status_code == 200
    assert data1["document"] == data2["document"]
    assert data1["labels"] == data2["labels"]


@pytest.mark.integration
@patch("aymurai.api.endpoints.routers.anonymizer.anonymizer.load_pipeline")
def test_should_return_prediction_without_storing_when_use_cache_false(
    mock_load_pipeline, client, db_session
):
    mock_pipeline = build_mock_pipeline()
    mock_load_pipeline.return_value = mock_pipeline

    text = "No cache storage text"
    response = client.post(
        "/anonymizer/predict",
        json={"text": text},
        params={"use_cache": False},
    )

    assert response.status_code == 200
    data = response.json()
    assert "document" in data
    assert "labels" in data

    paragraph_id = text_to_uuid(text)
    stored = db_session.get(AnonymizationParagraph, paragraph_id)
    assert stored is None


@pytest.mark.integration
def test_should_return_422_when_payload_is_invalid_json(client):
    response = client.post(
        "/anonymizer/predict",
        content="not json",
        headers={"Content-Type": "application/json"},
    )

    assert response.status_code == 422


@pytest.mark.integration
@patch("aymurai.api.endpoints.routers.anonymizer.anonymizer.load_pipeline")
def test_should_use_cache_by_default_when_param_omitted(
    mock_load_pipeline, client, db_session
):
    mock_pipeline = build_mock_pipeline()
    mock_load_pipeline.return_value = mock_pipeline

    text = "Default cache behavior"
    response = client.post(
        "/anonymizer/predict",
        json={"text": text},
    )

    assert response.status_code == 200

    paragraph_id = text_to_uuid(text)
    stored = db_session.get(AnonymizationParagraph, paragraph_id)
    assert stored is not None
    assert stored.text == text


@pytest.mark.integration
@patch("aymurai.api.endpoints.routers.anonymizer.anonymizer.load_pipeline")
def test_should_isolate_cache_when_different_texts(
    mock_load_pipeline, client, db_session
):
    mock_pipeline = build_mock_pipeline()
    mock_load_pipeline.return_value = mock_pipeline

    text1 = "First text for cache"
    text2 = "Second text for cache"

    para1_id = text_to_uuid(text1)
    para2_id = text_to_uuid(text2)

    labels1 = [build_label("PER", "Person1").model_dump(mode="json")]
    labels2 = [build_label("LOC", "Location1").model_dump(mode="json")]

    para1 = AnonymizationParagraph(id=para1_id, text=text1, prediction=labels1)
    para2 = AnonymizationParagraph(id=para2_id, text=text2, prediction=labels2)

    db_session.add_all([para1, para2])
    db_session.commit()

    response1 = client.post(
        "/anonymizer/predict",
        json={"text": text1},
        params={"use_cache": True},
    )

    response2 = client.post(
        "/anonymizer/predict",
        json={"text": text2},
        params={"use_cache": True},
    )

    assert response1.status_code == 200
    assert response2.status_code == 200

    data1 = response1.json()
    data2 = response2.json()

    assert data1["labels"] == labels1
    assert data2["labels"] == labels2
    assert data1["labels"] != data2["labels"]


@pytest.mark.integration
@patch("aymurai.api.endpoints.routers.anonymizer.anonymizer.load_pipeline")
def test_should_dedupe_duplicate_labels_when_returning_cached_prediction(
    mock_load_pipeline, client, db_session
):
    text = "EL SEÑOR JUEZ, doctor Tarte, señaló :"
    label = build_label("PER", "Tarte").model_dump(mode="json")
    label.update({"start_char": 22, "end_char": 27})
    label["attrs"].update(
        {
            "aymurai_alt_text": "Tarte",
            "aymurai_alt_start_char": 22,
            "aymurai_alt_end_char": 27,
            "aymurai_disambiguation": "fuzzy",
            "aymurai_anonymize": True,
        }
    )

    db_session.add(
        AnonymizationParagraph(
            id=text_to_uuid(text),
            text=text,
            prediction=[label, label],
        )
    )
    db_session.commit()

    response = client.post(
        "/anonymizer/predict",
        json={"text": text},
        params={"use_cache": True},
    )

    assert response.status_code == 200
    data = response.json()
    assert data["document"] == text
    assert data["labels"] == [label]
    mock_load_pipeline.assert_not_called()


@pytest.mark.integration
@patch("aymurai.api.endpoints.routers.anonymizer.anonymizer.load_pipeline")
def test_should_merge_cached_duplicate_labels_for_same_span_and_label(
    mock_load_pipeline, client, db_session
):
    text = (
        "Víctima: María Paula Trucha, DNI 23.456.789, quien se encuentra "
        "conectada con su cámara apagada."
    )
    dni_label = build_label("DNI", "23.456.789").model_dump(mode="json")
    dni_label.update({"start_char": 33, "end_char": 43})
    dni_label["attrs"].update(
        {
            "aymurai_alt_text": "23.456.789",
            "aymurai_alt_start_char": 33,
            "aymurai_alt_end_char": 43,
            "aymurai_label_instance": 2,
            "aymurai_disambiguation": "fuzzy",
            "aymurai_anonymize": True,
            "canonical_entity_id": "0bba6d15-1b0c-51f0-b2ca-4fdc8a57cb73",
        }
    )
    enriched_dni_label = {
        **dni_label,
        "attrs": {
            **dni_label["attrs"],
            "aymurai_label_subclass": ["23456789"],
        },
    }

    db_session.add(
        AnonymizationParagraph(
            id=text_to_uuid(text),
            text=text,
            prediction=[dni_label, enriched_dni_label],
        )
    )
    db_session.commit()

    response = client.post(
        "/anonymizer/predict",
        json={"text": text},
        params={"use_cache": True},
    )

    assert response.status_code == 200
    data = response.json()
    assert data["document"] == text
    assert data["labels"] == [enriched_dni_label]
    mock_load_pipeline.assert_not_called()


@pytest.mark.integration
@patch(
    "aymurai.api.endpoints.routers.anonymizer.anonymizer.map_canonical_entities_ner_preds"
)
@patch("aymurai.api.endpoints.routers.anonymizer.anonymizer.get_canonical_dates")
@patch("aymurai.api.endpoints.routers.anonymizer.anonymizer.build_canonical_entities")
def test_should_disambiguate_and_persist_paragraphs(
    mock_build_canonical_entities,
    mock_get_canonical_dates,
    mock_map_canonical_entities,
    client,
    db_session,
):
    mock_build_canonical_entities.return_value = []
    mock_get_canonical_dates.return_value = []
    mock_map_canonical_entities.side_effect = lambda predictions, canonical_entities: (
        predictions
    )

    text = "Ana Pérez denunció en el juzgado."
    body = {
        "paragraphs": [
            {
                "document": text,
                "labels": [build_label("PER", "Ana Pérez").model_dump(mode="json")],
            }
        ],
        "label_policies": {
            "PER": {"anonymize": True, "disambiguation": "none"},
        },
    }

    response = client.post("/anonymizer/disambiguate", json=body)

    assert response.status_code == 200
    payload = response.json()
    assert payload["label_policies"]["PER"]["disambiguation"] == "none"
    assert payload["data"][0]["labels"][0]["attrs"]["aymurai_disambiguation"] == "none"
    assert payload["data"][0]["labels"][0]["attrs"]["aymurai_anonymize"] is True

    paragraph_id = text_to_uuid(text)
    stored = db_session.get(AnonymizationParagraph, paragraph_id)
    assert stored is not None
    assert stored.prediction is not None
    assert stored.prediction[0]["text"] == "Ana Pérez"


@pytest.mark.integration
@patch(
    "aymurai.api.endpoints.routers.anonymizer.anonymizer.map_canonical_entities_ner_preds"
)
@patch("aymurai.api.endpoints.routers.anonymizer.anonymizer.get_canonical_dates")
@patch("aymurai.api.endpoints.routers.anonymizer.anonymizer.build_canonical_entities")
def test_should_dedupe_duplicate_labels_when_disambiguating_and_persisting(
    mock_build_canonical_entities,
    mock_get_canonical_dates,
    mock_map_canonical_entities,
    client,
    db_session,
):
    mock_build_canonical_entities.return_value = []
    mock_get_canonical_dates.return_value = []
    mock_map_canonical_entities.side_effect = lambda predictions, canonical_entities: (
        predictions
    )

    text = "EL SEÑOR JUEZ, doctor Tarte, señaló :"
    label = build_label("PER", "Tarte").model_dump(mode="json")
    label.update({"start_char": 22, "end_char": 27})
    label["attrs"].update(
        {
            "aymurai_alt_text": "Tarte",
            "aymurai_alt_start_char": 22,
            "aymurai_alt_end_char": 27,
        }
    )
    body = {
        "paragraphs": [{"document": text, "labels": [label, label]}],
        "label_policies": {
            "PER": {"anonymize": True, "disambiguation": "none"},
        },
    }

    response = client.post("/anonymizer/disambiguate", json=body)

    assert response.status_code == 200
    labels = response.json()["data"][0]["labels"]
    assert labels == [
        {
            **label,
            "attrs": {
                **label["attrs"],
                "aymurai_disambiguation": "none",
                "aymurai_anonymize": True,
            },
        }
    ]

    stored = db_session.get(AnonymizationParagraph, text_to_uuid(text))
    assert stored is not None
    assert stored.prediction is not None
    assert len(stored.prediction) == 1
    assert stored.prediction[0]["text"] == "Tarte"
    assert stored.prediction[0]["attrs"]["aymurai_disambiguation"] == "none"
    assert stored.prediction[0]["attrs"]["aymurai_anonymize"] is True


@pytest.mark.integration
def test_should_return_null_validation_when_paragraph_not_found(client):
    response = client.post(
        "/anonymizer/validation",
        json={"text": "Paragraph without validation"},
    )

    assert response.status_code == 200
    assert response.json() is None


@pytest.mark.integration
def test_should_return_validation_when_paragraph_exists(client, db_session):
    text = "Validated paragraph"
    labels = [build_label("PER", "María Soto").model_dump(mode="json")]
    db_session.add(
        AnonymizationParagraph(
            id=text_to_uuid(text),
            text=text,
            validation=labels,
        )
    )
    db_session.commit()

    response = client.post("/anonymizer/validation", json={"text": text})

    assert response.status_code == 200
    assert response.json() == labels


@pytest.mark.integration
@patch("aymurai.api.endpoints.routers.anonymizer.anonymizer.get_anonymizer")
def test_should_return_application_pdf_when_pdf_document_is_anonymized(
    mock_get_anonymizer,
    client,
    tmp_path,
):
    anonymized_path = _write_pdf(
        tmp_path / "output.pdf",
        lambda _doc, page: page.insert_text((72, 72), "Anonymized PDF output"),
    )
    mock_get_anonymizer.return_value = MagicMock(return_value=str(anonymized_path))

    annotations = {
        "data": [
            {
                "document": "Ana Perez presento el escrito",
                "labels": [build_label("PER", "Ana Perez").model_dump(mode="json")],
            }
        ],
        "label_policies": {"PER": {"anonymize": True, "disambiguation": "none"}},
        "render_policy": {"suffix_mode": "auto", "suffix_threshold": 1},
    }

    response = client.post(
        "/anonymizer/anonymize-document",
        data={"annotations": json.dumps(annotations)},
        files={"file": ("sample.pdf", b"%PDF-1.4 fake", "application/pdf")},
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert len(response.content) > 0


@patch("aymurai.api.endpoints.routers.anonymizer.anonymizer.subprocess.check_output")
@patch("aymurai.api.endpoints.routers.anonymizer.anonymizer.get_anonymizer")
def test_should_anonymize_document_when_annotations_are_valid(
    mock_get_anonymizer, mock_check_output, client, tmp_path
):
    # Fake anonymizer that writes a dummy docx output
    anonymized_path = str(tmp_path / "output.docx")
    with open(anonymized_path, "wb") as f:
        f.write(b"fake-docx-content")

    mock_anonymizer = MagicMock(return_value=anonymized_path)
    mock_get_anonymizer.return_value = mock_anonymizer

    def fake_convert(*args, **kwargs):
        cmd = args[0]
        source_path = cmd[-1]
        output_path = source_path.rsplit(".", 1)[0] + ".odt"
        with open(output_path, "wb") as output_file:
            output_file.write(b"odt-content")
        return "ok"

    mock_check_output.side_effect = fake_convert
    annotations = {
        "data": [
            {
                "document": "Ana Pérez denunció en el juzgado.",
                "labels": [build_label("PER", "Ana Pérez").model_dump(mode="json")],
            }
        ],
        "label_policies": {"PER": {"anonymize": True, "disambiguation": "fuzzy"}},
        "render_policy": {"suffix_mode": "auto", "suffix_threshold": 1},
    }

    response = client.post(
        "/anonymizer/anonymize-document",
        data={"annotations": json.dumps(annotations)},
        files={
            "file": (
                "sample.docx",
                b"input-document",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/octet-stream"
    assert len(response.content) > 0


@pytest.mark.integration
@patch("aymurai.api.endpoints.routers.anonymizer.anonymizer.subprocess.check_output")
@patch("aymurai.api.endpoints.routers.anonymizer.anonymizer.get_anonymizer")
def test_should_exclude_null_alt_attrs_from_anonymize_document_preds(
    mock_get_anonymizer, mock_check_output, client, tmp_path
):
    anonymized_path = str(tmp_path / "output.docx")
    with open(anonymized_path, "wb") as f:
        f.write(b"fake-docx-content")

    mock_anonymizer = MagicMock(return_value=anonymized_path)
    mock_get_anonymizer.return_value = mock_anonymizer

    def fake_convert(*args, **kwargs):
        cmd = args[0]
        source_path = cmd[-1]
        output_path = source_path.rsplit(".", 1)[0] + ".odt"
        with open(output_path, "wb") as output_file:
            output_file.write(b"odt-content")
        return "ok"

    mock_check_output.side_effect = fake_convert
    annotations = {
        "data": [
            {
                "document": "Ana Perez denuncio en el juzgado.",
                "labels": [build_label("PER", "Ana Perez").model_dump(mode="json")],
            }
        ],
        "label_policies": {"PER": {"anonymize": True, "disambiguation": "fuzzy"}},
        "render_policy": {"suffix_mode": "auto", "suffix_threshold": 1},
    }

    response = client.post(
        "/anonymizer/anonymize-document",
        data={"annotations": json.dumps(annotations)},
        files={
            "file": (
                "sample.docx",
                b"input-document",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200
    preds = mock_anonymizer.call_args[0][1]
    assert preds[0]["labels"][0]["text"] == "Ana Perez"

    attrs = preds[0]["labels"][0]["attrs"]
    assert "aymurai_alt_text" not in attrs
    assert "aymurai_alt_start_char" not in attrs
    assert "aymurai_alt_end_char" not in attrs


@pytest.mark.integration
@patch("aymurai.api.endpoints.routers.anonymizer.anonymizer.subprocess.check_output")
@patch("aymurai.api.endpoints.routers.anonymizer.anonymizer.get_anonymizer")
def test_should_return_500_when_anonymize_document_conversion_fails(
    mock_get_anonymizer, mock_check_output, client, tmp_path
):
    # Fake anonymizer that writes a dummy output
    anonymized_path = str(tmp_path / "output.docx")
    with open(anonymized_path, "wb") as f:
        f.write(b"fake-docx-content")

    mock_anonymizer = MagicMock(return_value=anonymized_path)
    mock_get_anonymizer.return_value = mock_anonymizer

    mock_check_output.side_effect = subprocess.CalledProcessError(
        1,
        ["libreoffice"],
        output=b"conversion failed",
    )
    annotations = {
        "data": [{"document": "text", "labels": []}],
        "label_policies": None,
        "render_policy": {"suffix_mode": "auto", "suffix_threshold": 1},
    }

    response = client.post(
        "/anonymizer/anonymize-document",
        data={"annotations": json.dumps(annotations)},
        files={
            "file": (
                "sample.docx",
                b"input-document",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 500
